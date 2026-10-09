# -*- coding: utf-8 -*-
"""
Multi-format attachment parser for Super-Harnes.
Supports extracting content from:
- Plain text / Markdown / Source code (.txt, .md, .py, .json, etc.)
- Word documents (.docx)
- PDF documents (.pdf)
- Excel spreadsheets (.xlsx, .xls, .csv)
- Images (.png, .jpg, .jpeg, .webp, .gif, .svg)
"""

import os
import base64
import logging
from pathlib import Path
from typing import Dict, Any, Optional

logger = logging.getLogger("attachment_parser")

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg", ".bmp"}
DOC_EXTENSIONS = {".docx", ".doc"}
PDF_EXTENSIONS = {".pdf"}
EXCEL_EXTENSIONS = {".xlsx", ".xls", ".csv"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".json", ".yaml", ".yml", ".py", ".js", ".ts", ".jsx", ".tsx",
    ".html", ".css", ".scss", ".xml", ".sql", ".sh", ".bat", ".log", ".ini", ".conf", ".toml"
}

def detect_attachment_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in IMAGE_EXTENSIONS:
        return "image"
    elif ext in PDF_EXTENSIONS:
        return "pdf"
    elif ext in DOC_EXTENSIONS:
        return "word"
    elif ext in EXCEL_EXTENSIONS:
        return "excel"
    elif ext in TEXT_EXTENSIONS:
        return "text"
    return "other"

def parse_text_file(path: Path, max_chars: int = 30000) -> str:
    for enc in ("utf-8", "gb18030", "gbk", "latin-1"):
        try:
            with open(path, "r", encoding=enc, errors="replace") as f:
                content = f.read(max_chars + 100)
                if len(content) > max_chars:
                    return content[:max_chars] + f"\n\n...[文件过长，已截取前 {max_chars} 字符]..."
                return content
        except Exception:
            continue
    return "[无法使用常见编码读取此文本文件]"

def parse_pdf_file(path: Path, max_pages: int = 30, max_chars: int = 30000) -> str:
    try:
        import pymupdf
        doc = pymupdf.open(str(path))
        total_pages = len(doc)
        extracted = [f"【PDF 文件元信息: 共 {total_pages} 页】"]
        
        cur_chars = 0
        for page_idx in range(min(total_pages, max_pages)):
            page = doc[page_idx]
            text = page.get_text()
            if text.strip():
                extracted.append(f"\n--- [第 {page_idx + 1} 页] ---\n{text.strip()}")
                cur_chars += len(text)
                if cur_chars >= max_chars:
                    extracted.append(f"\n...[PDF 内容已截取前 {page_idx + 1} 页 (约 {cur_chars} 字符)]...")
                    break
                    
        doc.close()
        return "\n".join(extracted)
    except Exception as e:
        logger.warning(f"Failed to parse PDF {path}: {e}")
        return f"[PDF 解析异常: {e}]"

def parse_docx_file(path: Path, max_chars: int = 30000) -> str:
    try:
        import docx
        doc = docx.Document(str(path))
        lines = []
        
        # 提取标题与段落
        for p in doc.paragraphs:
            if p.text.strip():
                lines.append(p.text.strip())
                
        # 提取表格
        if doc.tables:
            lines.append(f"\n【文档包含 {len(doc.tables)} 个表格】:")
            for t_idx, tbl in enumerate(doc.tables):
                lines.append(f"\n[表格 {t_idx + 1}]:")
                for row in tbl.rows[:30]: # 最多前 30 行
                    row_data = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    lines.append("| " + " | ".join(row_data) + " |")
                    
        full_text = "\n".join(lines)
        if len(full_text) > max_chars:
            return full_text[:max_chars] + f"\n\n...[Word 文档内容已截取前 {max_chars} 字符]..."
        return full_text if full_text.strip() else "[Word 文档内容为空]"
    except Exception as e:
        logger.warning(f"Failed to parse docx {path}: {e}")
        return f"[Word 文档解析异常: {e}]"

def _df_to_markdown_native(df, max_rows: int = 50) -> str:
    if df.empty:
        return "[空表格数据]"
    cols = [str(c) for c in df.columns]
    header = "| " + " | ".join(cols) + " |"
    separator = "| " + " | ".join(["---"] * len(cols)) + " |"
    rows = []
    for _, row in df.head(max_rows).iterrows():
        row_str = [str(val).replace("\n", " ").replace("|", "\\|") for val in row]
        rows.append("| " + " | ".join(row_str) + " |")
    return "\n".join([header, separator] + rows)

def parse_excel_file(path: Path, max_rows: int = 50) -> str:
    ext = path.suffix.lower()
    if ext == ".csv":
        try:
            import csv
            lines = []
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                rows = []
                for idx, r in enumerate(reader):
                    if idx >= max_rows:
                        break
                    rows.append(r)
            if rows:
                cols = rows[0]
                lines.append(f"【CSV 数据预览 (前 {len(rows)} 行，共 {len(cols)} 列)】:")
                header = "| " + " | ".join([str(c) for c in cols]) + " |"
                sep = "| " + " | ".join(["---"] * len(cols)) + " |"
                lines.append(header)
                lines.append(sep)
                for r in rows[1:]:
                    r_str = [str(val).replace("\n", " ").replace("|", "\\|") for val in r]
                    lines.append("| " + " | ".join(r_str) + " |")
                return "\n".join(lines)
            return "[空 CSV 数据]"
        except Exception as e:
            logger.warning(f"Failed to parse csv {path}: {e}")

    try:
        import pandas as pd
        lines = []
        xl = pd.ExcelFile(str(path))
        sheet_names = xl.sheet_names
        lines.append(f"【Excel 包含工作表: {', '.join(sheet_names)}】")
        for sheet in sheet_names[:5]:
            df = xl.parse(sheet, nrows=max_rows)
            lines.append(f"\n--- [工作表: {sheet}] (前 {len(df)} 行，{len(df.columns)} 列) ---")
            lines.append(_df_to_markdown_native(df, max_rows=max_rows))
        return "\n".join(lines)
    except ImportError:
        return "[Excel 表格解析提示: 未安装 pandas/openpyxl 依赖，可在终端执行 pip install pandas openpyxl 开启原生表格预览]"
    except Exception as e:
        logger.warning(f"Failed to parse excel {path}: {e}")
        return f"[表格解析异常: {e}]"

def parse_image_file(path: Path) -> Dict[str, Any]:
    try:
        with open(path, "rb") as f:
            b64 = base64.b64encode(f.read()).decode("utf-8")
        ext = path.suffix.lower().lstrip(".")
        mime = f"image/{ext}" if ext != "jpg" else "image/jpeg"
        return {
            "mime": mime,
            "data_url": f"data:{mime};base64,{b64}",
            "size": path.stat().st_size
        }
    except Exception as e:
        logger.warning(f"Failed to read image {path}: {e}")
        return {}

def extract_attachment_summary(file_path: Path, filename: Optional[str] = None) -> Dict[str, Any]:
    """提取附件内容结构化摘要，用于注入 Agent 上下文"""
    name = filename or file_path.name
    ftype = detect_attachment_type(name)
    size_bytes = file_path.stat().st_size if file_path.exists() else 0
    
    extracted_text = ""
    image_meta = None
    
    if ftype == "text":
        extracted_text = parse_text_file(file_path)
    elif ftype == "pdf":
        extracted_text = parse_pdf_file(file_path)
    elif ftype == "word":
        extracted_text = parse_docx_file(file_path)
    elif ftype == "excel":
        extracted_text = parse_excel_file(file_path)
    elif ftype == "image":
        image_meta = parse_image_file(file_path)
        extracted_text = f"[用户上传了图片: {name} (大小: {size_bytes} 字节)，图片已保存在: {file_path.as_posix()}]"
    else:
        extracted_text = f"[其他类型文件: {name} (大小: {size_bytes} 字节)，保存在: {file_path.as_posix()}]"

    return {
        "name": name,
        "type": ftype,
        "size": size_bytes,
        "path": file_path.as_posix(),
        "extracted_text": extracted_text,
        "image_meta": image_meta
    }
